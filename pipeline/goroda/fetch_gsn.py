# Реестр объектов капитального строительства, поднадзорных Инспекции ГСН РТ (gsn.tatarstan.ru, еженедельно, xlsx).
# Скачивает самый свежий файл со страницы реестра → gsn/registry_<дата>.xlsx и печатает структуру. Госсайт РТ — curl --noproxy.
import re, html, subprocess, pathlib

ROOT = pathlib.Path(__file__).parent / 'gsn'
PAGE = 'https://gsn.tatarstan.ru/reestr-obektov-kapitalnogo-stroitelstva-i.htm'
UA = 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 Chrome/126.0.0.0 Safari/537.36'
curl = lambda url, out=None: subprocess.run(['curl', '-s', '-k', '--noproxy', '*', '--max-time', '120', '-A', UA, url] + (['-o', str(out)] if out else []),
                                            capture_output=True).stdout

def latest():
    s = curl(PAGE).decode('utf-8', 'replace')
    # публикация: дата рядом со ссылкой на pub_*.xlsx; берём самую позднюю дату
    items = []
    for m in re.finditer(r'(/file/pub/pub_\d+\.xlsx)', s):
        ctx = html.unescape(re.sub(r'<[^>]+>', ' ', s[max(0, m.start() - 1500):m.start()]))
        d = re.findall(r'(\d\d)\.(\d\d)\.(20\d\d)', ctx)
        if d: items.append(('{2}-{1}-{0}'.format(*d[-1]), 'https://gsn.tatarstan.ru' + m.group(1)))
    return max(items) if items else None

if __name__ == '__main__':
    ROOT.mkdir(exist_ok=True)
    date, url = latest()
    out = ROOT / f'registry_{date}.xlsx'
    if not out.exists(): curl(url, out)
    print(date, url, out.stat().st_size, 'байт')
    import openpyxl
    wb = openpyxl.load_workbook(out, read_only=True)
    for ws in wb.worksheets:
        print('лист', ws.title, ws.max_row, 'строк')
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i < 6: print('  ', [str(c)[:60] if c is not None else '' for c in row][:10])
            else: break
